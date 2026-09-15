import os
import json
import time
import urllib.request
import urllib.parse
import ssl
import openpyxl

def load_env():
    env_vars = {}
    if os.path.exists('.env'):
        with open('.env', 'r', encoding='utf-8') as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith('#') and '=' in line:
                    k, v = line.split('=', 1)
                    env_vars[k.strip()] = v.strip()
    return env_vars

env = load_env()
CLIENT_ID = env.get('NEXT_PUBLIC_NAVER_MAP_CLIENT_ID')
CLIENT_SECRET = env.get('NAVER_MAP_CLIENT_SECRET')

if not CLIENT_ID or not CLIENT_SECRET:
    raise ValueError("Missing Naver API credentials in .env")

# Geocoding Cache Setup
CACHE_FILE = r'raw_data\geocode_cache.json'
os.makedirs('raw_data', exist_ok=True)
os.makedirs(r'public\data', exist_ok=True)

geocode_cache = {}
if os.path.exists(CACHE_FILE):
    try:
        with open(CACHE_FILE, 'r', encoding='utf-8') as f:
            geocode_cache = json.load(f)
        print(f"Loaded {len(geocode_cache)} cached geocode entries.")
    except Exception as e:
        print(f"Cache load error: {e}")

ssl_ctx = ssl.create_default_context()
ssl_ctx.check_hostname = False
ssl_ctx.verify_mode = ssl.CERT_NONE

def geocode_address(address, town=''):
    clean_addr = address.strip()
    if clean_addr in geocode_cache and geocode_cache[clean_addr] is not None:
        return geocode_cache[clean_addr]
    
    encoded = urllib.parse.quote(clean_addr)
    url = f"https://maps.apigw.ntruss.com/map-geocode/v2/geocode?query={encoded}"
    req = urllib.request.Request(url, headers={
        'x-ncp-apigw-api-key-id': CLIENT_ID,
        'x-ncp-apigw-api-key': CLIENT_SECRET,
        'Accept': 'application/json'
    })

    try:
        with urllib.request.urlopen(req, context=ssl_ctx, timeout=8) as res:
            data = json.loads(res.read().decode('utf-8'))
            addrs = data.get('addresses', [])
            if addrs:
                first = addrs[0]
                result = {
                    'lat': float(first.get('y')),
                    'lng': float(first.get('x')),
                    'roadAddress': first.get('roadAddress', clean_addr),
                    'jibunAddress': first.get('jibunAddress', '')
                }
                geocode_cache[clean_addr] = result
                return result
    except Exception as e:
        print(f"Geocoding error for '{clean_addr}': {e}")
    
    # Fallback
    if town and town not in clean_addr:
        fallback_query = f"전라남도 함평군 {town} {clean_addr}"
        try:
            encoded_fallback = urllib.parse.quote(fallback_query)
            url_fallback = f"https://maps.apigw.ntruss.com/map-geocode/v2/geocode?query={encoded_fallback}"
            req_fb = urllib.request.Request(url_fallback, headers={
                'x-ncp-apigw-api-key-id': CLIENT_ID,
                'x-ncp-apigw-api-key': CLIENT_SECRET,
                'Accept': 'application/json'
            })
            with urllib.request.urlopen(req_fb, context=ssl_ctx, timeout=8) as res:
                data = json.loads(res.read().decode('utf-8'))
                addrs = data.get('addresses', [])
                if addrs:
                    first = addrs[0]
                    result = {
                        'lat': float(first.get('y')),
                        'lng': float(first.get('x')),
                        'roadAddress': first.get('roadAddress', clean_addr),
                        'jibunAddress': first.get('jibunAddress', '')
                    }
                    geocode_cache[clean_addr] = result
                    return result
        except Exception:
            pass

    geocode_cache[clean_addr] = None
    return None

def determine_category(name, orig_cat):
    s_name = str(name).strip()
    o_cat = str(orig_cat).strip()

    if any(k in s_name for k in ['하나로마트', '농협 하나로마트', '로컬푸드']):
        return '농협·하나로마트'
    if any(k in s_name for k in ['주유소', '충전소', 'LPG']):
        return '주유소·충전소'
    if any(k in s_name for k in ['약국']):
        return '약국'
    if any(k in s_name for k in ['의원', '병원', '치과', '한의원']):
        return '병원·의원'
    if any(k in s_name for k in ['카페', '커피', '베이커리', '제과', '디저트', '떡']) or o_cat == '카페/베이커리':
        return '카페·디저트'
    if any(k in s_name for k in ['식당', '반점', '통닭', '치킨', '국밥', '비빔밥', '갈비', '회관', '추어탕', '짜장', '분식', '피자', '호프', '가든']) or o_cat == '음식점업':
        return '음식점'
    if any(k in s_name for k in ['마트', '슈퍼', '편의점', '유통', '상회', '식자재']) or o_cat == '소매업':
        if any(k in s_name for k in ['정육', '축산', '수산', '수협', '식육']):
            return '정육·수산·식품'
        return '마트·편의점'
    if any(k in s_name for k in ['미용', '헤어', '이용원', '이발', '네일']):
        return '미용·뷰티'
    if any(k in s_name for k in ['세탁', '크리닝', '목욕', '사우나', '인테리어', '사진관']):
        return '생활·서비스'
    if any(k in s_name for k in ['학원', '교습', '문구', '서점']):
        return '학원·교육'
    if any(k in s_name for k in ['모텔', '펜션', '여관', '민박', '호텔']):
        return '숙박'
    if any(k in s_name for k in ['정비', '카센타', '타이어', '모터스', '공업사']):
        return '자동차·정비'
    if any(k in s_name for k in ['비료', '종묘', '농자재', '기계', '철물']):
        return '농자재·철물'

    if o_cat == '음식점업': return '음식점'
    if o_cat == '소매업': return '마트·편의점'
    if o_cat == '개인서비스업': return '생활·서비스'
    if o_cat == '제조업': return '기타'
    if o_cat == '보건업': return '병원·의원'
    return '기타'

def evaluate_status(name, town, standard_cat, address):
    s_name = str(name).strip()

    # 1. 함평읍 소재 대형 하나로마트 / 농협본점 (30억 초과)
    if '하나로마트' in s_name or '축협' in s_name:
        if town == '함평읍':
            return {
                'status': 'unavailable',
                'badge': '사용불가',
                'reason': '연매출 30억 원 초과 제한 사업장 (면 지역이 아닌 함평읍 본점은 선불카드 사용 불가)',
                'rule_type': 'EXCEED_REVENUE'
            }
        else:
            return {
                'status': 'available',
                'badge': '예외가능',
                'reason': '면 지역 하나로마트 예외 허용 대상 (민생지원금 선불카드 사용 가능)',
                'rule_type': 'MYEON_HANARO_EXCEPTION'
            }

    # 2. 대형 식자재마트
    if any(k in s_name for k in ['식자재마트', '나비식자재']):
        return {
            'status': 'verify',
            'badge': '확인필요',
            'reason': '연매출 30억 원 기준 초과 여부 확인 필요 (대형 매장으로 결제 전 가맹점 사전 문의 권장)',
            'rule_type': 'VERIFY_REVENUE'
        }

    # 3. 유흥 / 사행성
    if any(k in s_name for k in ['단란주점', '유흥주점', '룸싸롱']):
        return {
            'status': 'unavailable',
            'badge': '사용불가',
            'reason': '유흥·사행성 업종 제한 (선불카드 사용 불가)',
            'rule_type': 'RESTRICTED_INDUSTRY'
        }

    # 4. 프랜차이즈 직영점
    if any(k in s_name for k in ['스타벅스', '직영점']):
        return {
            'status': 'unavailable',
            'badge': '사용불가',
            'reason': '대기업 본사 직영점 제한 대상 (선불카드 사용 불가)',
            'rule_type': 'DIRECT_FRANCHISE'
        }

    # 5. 일반 등록 가맹점 -> 사용 가능
    return {
        'status': 'available',
        'badge': '사용가능',
        'reason': '함평군 2026년 민생지원금 사용 가능 가맹점',
        'rule_type': 'STANDARD_MERCHANT'
    }

print("Loading workbook raw_data/hplovegiftcard_2026.xlsx ...")
wb = openpyxl.load_workbook(r'raw_data\hplovegiftcard_2026.xlsx', data_only=True)
ws = wb['★총괄']

raw_rows = []
for r in range(4, ws.max_row + 1):
    num = ws.cell(r, 1).value
    name = ws.cell(r, 2).value
    addr = ws.cell(r, 3).value
    town = ws.cell(r, 4).value
    orig_cat = ws.cell(r, 5).value
    reg_date = ws.cell(r, 6).value
    if name and addr:
        raw_rows.append({
            'num': num,
            'name': str(name).strip(),
            'address': str(addr).strip(),
            'town': str(town).strip() if town else '',
            'orig_cat': str(orig_cat).strip() if orig_cat else '기타',
            'reg_date': str(reg_date).strip() if reg_date else ''
        })

# Explicitly add major restricted places that residents frequently ask about
EXPLICIT_RESTRICTED = [
    {
        'num': 9901,
        'name': '함평농협 하나로마트 본점',
        'address': '전라남도 함평군 함평읍 중앙길 78',
        'town': '함평읍',
        'orig_cat': '농축협본점',
        'reg_date': '2026.09 (공식지침제한)'
    },
    {
        'num': 9902,
        'name': '함평축협 하나로마트 본점',
        'address': '전라남도 함평군 함평읍 서부길 57',
        'town': '함평읍',
        'orig_cat': '농축협본점',
        'reg_date': '2026.09 (공식지침제한)'
    },
    {
        'num': 9903,
        'name': '함평축협 한우프라자 본점',
        'address': '전라남도 함평군 함평읍 서부길 57',
        'town': '함평읍',
        'orig_cat': '음식점업',
        'reg_date': '2026.09 (공식지침제한)'
    }
]

raw_rows.extend(EXPLICIT_RESTRICTED)
print(f"Total stores to process (including explicit restricted): {len(raw_rows)}")

processed_stores = []
stats = {'available': 0, 'unavailable': 0, 'verify': 0, 'geocoded': 0, 'failed_geocode': 0}
categories_count = {}

for idx, item in enumerate(raw_rows):
    std_cat = determine_category(item['name'], item['orig_cat'])
    rule = evaluate_status(item['name'], item['town'], std_cat, item['address'])
    
    geo = geocode_address(item['address'], item['town'])
    if not geo:
        parts = item['address'].split()
        if len(parts) >= 4:
            simplified = ' '.join(parts[:4])
            geo = geocode_address(simplified, item['town'])

    stats[rule['status']] += 1
    categories_count[std_cat] = categories_count.get(std_cat, 0) + 1

    lat = geo['lat'] if geo else None
    lng = geo['lng'] if geo else None

    if geo:
        stats['geocoded'] += 1
    else:
        stats['failed_geocode'] += 1

    store_entry = {
        'id': f"hp_{idx + 1:04d}",
        'name': item['name'],
        'address': item['address'],
        'roadAddress': geo['roadAddress'] if geo else item['address'],
        'town': item['town'],
        'category': std_cat,
        'origCategory': item['orig_cat'],
        'status': rule['status'],
        'badge': rule['badge'],
        'reason': rule['reason'],
        'ruleType': rule['rule_type'],
        'lat': lat,
        'lng': lng,
        'regDate': item['reg_date'],
        'isGasStation': std_cat == '주유소·충전소'
    }
    processed_stores.append(store_entry)

# Save final stores.json
final_path = r'public\data\stores.json'
with open(final_path, 'w', encoding='utf-8') as f:
    json.dump({
        'metadata': {
            'title': '함평콕 2026 민생회복지원금 사용처 마스터 데이터',
            'totalCount': len(processed_stores),
            'updatedAt': '2026-09-15',
            'source': '함평군청 공식 함평사랑상품권 가맹점 현황 (2026-03) + 2026 민생지원금 제한지침',
            'stats': stats
        },
        'categories': sorted(list(categories_count.keys())),
        'towns': ['함평읍', '손불면', '신광면', '학교면', '엄다면', '대동면', '나산면', '해보면', '월야면'],
        'stores': processed_stores
    }, f, ensure_ascii=False, indent=2)

with open(CACHE_FILE, 'w', encoding='utf-8') as f:
    json.dump(geocode_cache, f, ensure_ascii=False, indent=2)

print("\n--- PROCESSING COMPLETED ---")
print(f"Saved {len(processed_stores)} stores to {final_path}")
print(f"Stats: {stats}")
